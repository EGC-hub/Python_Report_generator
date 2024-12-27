import comtypes.client
from docx import Document
import os

def get_versioned_filename(base_name, extension, directory):
    """Generate a filename with an incremented version number in the specified directory."""
    counter = 1
    while True:
        filename = f"{base_name}_v{counter}.{extension}"
        full_path = os.path.join(directory, filename)
        if not os.path.exists(full_path):
            return full_path
        counter += 1

def replace_placeholders(doc_path, output_path, data):
    doc = Document(doc_path)

    # Replace placeholders in paragraphs
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(key, value)

    # Save the updated Word document
    doc.save(output_path)

def convert_to_pdf(word_path, pdf_path):
    word = comtypes.client.CreateObject("Word.Application")
    word.Visible = False  # Hide Word application
    try:
        doc = word.Documents.Open(word_path)
        doc.SaveAs(pdf_path, FileFormat=17)  # 17 represents wdFormatPDF
        doc.Close()
    finally:
        word.Quit()

if __name__ == '__main__':
    data = {
        '[Report Title]': 'Report 3',
        '[Your Name or Organization]': 'Euro Global Consultancy',
        '[Insert Date]': '27/12/2024',
        '[Your Email Address]': 'sadullah@mail.com',
        '[Your Phone Number]': '1234567890',
        '[Your Website]': 'euroglobalconsultancy.com'
    }

    # Get absolute paths
    base_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(base_dir, 'template.docx')

    # Generate versioned filenames in the current directory
    word_output_path = get_versioned_filename("filled", "docx", base_dir)
    pdf_output_path = get_versioned_filename("filled", "pdf", base_dir)

    # Process and generate the outputs
    replace_placeholders(template_path, word_output_path, data)
    convert_to_pdf(word_output_path, pdf_output_path)

    print(f'Report generated:\nWord file: {word_output_path}\nPDF file: {pdf_output_path}')