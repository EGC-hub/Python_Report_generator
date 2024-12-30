import mysql.connector
from mysql.connector import Error
from docx import Document
import comtypes.client
import os

def connect_to_database(host, user, password, database):
    """
    Connect to a MySQL database.
    """
    try:
        connection = mysql.connector.connect(
            host=host,
            user=user,
            password=password,
            database=database
        )
        if connection.is_connected():
            print(f"Connected to MySQL database '{database}'")
            return connection
    except Error as e:
        print(f"Error connecting to MySQL: {e}")
        return None

def fetch_data_from_database(connection, record_id):
    """
    Fetch data from the database for the given record ID.
    """
    try:
        cursor = connection.cursor(dictionary=True)
        query = "SELECT * FROM reports WHERE id = %s"
        cursor.execute(query, (record_id,))
        result = cursor.fetchone()
        return result
    except Error as e:
        print(f"Error fetching data from database: {e}")
        return None
    finally:
        if cursor:
            cursor.close()

def get_versioned_filename(base_name, extension, directory):
    """Generate a filename with an incremented version number in the specified directory."""
    counter = 1
    while True:
        filename = f"{base_name}_v{counter}.{extension}"
        full_path = os.path.join(directory, filename)
        if not os.path.exists(full_path):
            return full_path
        counter += 1

def replace_placeholders(doc_path, output_path, data):
    doc = Document(doc_path)

    # Replace placeholders in paragraphs
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(key, value)

    # Save the updated Word document
    doc.save(output_path)

def convert_to_pdf(word_path, pdf_path):
    word = comtypes.client.CreateObject("Word.Application")
    word.Visible = False  # Hide Word application
    try:
        doc = word.Documents.Open(word_path)
        doc.SaveAs(pdf_path, FileFormat=17)  # 17 represents wdFormatPDF
        doc.Close()
    finally:
        word.Quit()

if __name__ == '__main__':
    # Configuration for database connection
    db_config = {
        "host": "localhost",
        "user": "root",
        "password": "",
        "database": "python_database"
    }

    # Connect to the database
    connection = connect_to_database(**db_config)
    if not connection:
        exit("Failed to connect to the database.")

    # Fetch data for the specific record ID
    record_id = 9  # Change this ID as needed
    record = fetch_data_from_database(connection, record_id)
    if not record:
        exit(f"No record found with ID {record_id}.")

    # Disconnect from the database
    connection.close()

    # Map database fields to placeholders
    data = {
        '[Report Title]': record.get('report_title', 'N/A'),
        '[Your Name or Organization]': record.get('organization_name', 'N/A'),
        '[Insert Date]': record.get('report_date', 'N/A'),
        '[Your Email Address]': record.get('email_address', 'N/A'),
        '[Your Phone Number]': record.get('phone_number', 'N/A'),
        '[Your Website]': record.get('website', 'N/A')
    }

    # Get absolute paths
    base_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(base_dir, 'template.docx')

    # Generate versioned filenames in the current directory
    word_output_path = get_versioned_filename("filled", "docx", base_dir)
    pdf_output_path = get_versioned_filename("filled", "pdf", base_dir)

    # Process and generate the outputs
    replace_placeholders(template_path, word_output_path, data)
    convert_to_pdf(word_output_path, pdf_output_path)

    print(f"Report generated:\nWord file: {word_output_path}\nPDF file: {pdf_output_path}")